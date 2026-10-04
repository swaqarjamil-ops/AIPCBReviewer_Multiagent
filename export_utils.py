"""
export_utils.py
---------------
Builds downloadable MS Word (.docx) and PDF versions of the final
markdown report. The functions return in-memory bytes so Streamlit can
serve them directly without creating temporary files.
"""

import io
import re
from datetime import datetime
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def _generated_on() -> str:
    """Return today's date formatted for the report header, e.g. 'October 04, 2026'."""
    return datetime.now().strftime("%B %d, %Y")


def _plain(text: str) -> str:
    """Remove common markdown emphasis/link syntax for document output."""
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = re.sub(r"[*_`]+", "", text)
    return text.strip()


def _is_table_separator(line: str) -> bool:
    """Return True for a markdown table separator row."""
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def _parse_markdown(report: str):
    """Yield simple typed blocks from the report markdown."""
    lines = report.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            yield ("heading", len(heading.group(1)), _plain(heading.group(2)))
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and _is_table_separator(lines[i + 1]):
            headers = [_plain(c.strip()) for c in line.strip().strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([_plain(c.strip()) for c in lines[i].strip().strip("|").split("|")])
                i += 1
            yield ("table", headers, rows)
            continue

        if re.match(r"^[-*]\s+", line):
            yield ("bullet", _plain(re.sub(r"^[-*]\s+", "", line)))
            i += 1
            continue

        if re.match(r"^\d+[.)]\s+", line):
            yield ("numbered", _plain(re.sub(r"^\d+[.)]\s+", "", line)))
            i += 1
            continue

        paragraph = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or re.match(r"^#{1,3}\s+", nxt)
                or nxt.startswith("|")
                or re.match(r"^[-*]\s+", nxt)
                or re.match(r"^\d+[.)]\s+", nxt)
            ):
                break
            paragraph.append(nxt)
            i += 1
        yield ("paragraph", _plain(" ".join(paragraph)))


def report_to_docx(report: str, title: str = "CircuitMind Enterprise AI Engineering Report") -> bytes:
    """Convert markdown report to a DOCX document and return its bytes."""
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(18)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date_p.add_run(f"Generated on {_generated_on()}")
    date_run.italic = True
    date_run.font.size = Pt(10)
    date_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    for block in _parse_markdown(report):
        kind = block[0]
        if kind == "heading":
            _, level, text = block
            p = doc.add_heading(text, level=min(level + 1, 3))
        elif kind == "paragraph":
            doc.add_paragraph(block[1])
        elif kind == "bullet":
            doc.add_paragraph(block[1], style="List Bullet")
        elif kind == "numbered":
            doc.add_paragraph(block[1], style="List Number")
        elif kind == "table":
            _, headers, rows = block
            table = doc.add_table(rows=1, cols=max(1, len(headers)))
            table.style = "Table Grid"
            for j, value in enumerate(headers):
                table.rows[0].cells[j].text = value
            for row in rows:
                cells = table.add_row().cells
                for j in range(min(len(cells), len(row))):
                    cells[j].text = row[j]

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def report_to_pdf(report: str, title: str = "CircuitMind Enterprise AI  Engineering Report") -> bytes:
    """Convert markdown report to an A4 PDF and return its bytes."""
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontSize=8.5, leading=11))
    styles.add(ParagraphStyle(name="ReportH1", parent=styles["Heading1"], fontSize=15, leading=18, spaceAfter=7))
    styles.add(ParagraphStyle(name="ReportH2", parent=styles["Heading2"], fontSize=12, leading=15, spaceBefore=8, spaceAfter=5))
    styles.add(ParagraphStyle(name="ReportBullet", parent=styles["BodyText"], fontSize=8.5, leading=11, leftIndent=12, bulletIndent=3))
    styles.add(ParagraphStyle(name="ReportDate", parent=styles["BodyText"], fontSize=9.5, leading=12, alignment=1, textColor=colors.grey, spaceAfter=6))

    story = [
        Paragraph(escape(title), styles["Title"]),
        Paragraph(f"Generated on {escape(_generated_on())}", styles["ReportDate"]),
        Spacer(1, 8),
    ]

    for block in _parse_markdown(report):
        kind = block[0]
        if kind == "heading":
            _, level, text = block
            style = styles["ReportH1"] if level == 1 else styles["ReportH2"]
            story.append(Paragraph(escape(text), style))
        elif kind == "paragraph":
            story.append(Paragraph(escape(block[1]), styles["ReportBody"]))
            story.append(Spacer(1, 4))
        elif kind == "bullet":
            story.append(Paragraph(escape(block[1]), styles["ReportBullet"], bulletText="•"))
        elif kind == "numbered":
            story.append(Paragraph(escape(block[1]), styles["ReportBullet"]))
        elif kind == "table":
            _, headers, rows = block
            data = [[Paragraph(escape(h), styles["ReportBody"]) for h in headers]]
            for row in rows:
                data.append([Paragraph(escape(c), styles["ReportBody"]) for c in row])
            table = Table(data, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(table)
            story.append(Spacer(1, 6))

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=title,
    )
    pdf.build(story)
    return output.getvalue()
